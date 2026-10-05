"""Lee el Excel "TURNOS Y CANCHAS" y lo convierte al texto de horario de la app.

Cada hoja (HOMBRES / MUJERES) tiene una tabla por fecha. Se usa la última fecha
con partidos en ambas hojas. Los equipos vienen abreviados ("tortu", "eq banco"),
así que se cruzan con los nombres reales de los clubes del masivo ("Club: …").
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from rapidfuzz import fuzz

from processor.names import normalize_name

SHEET_CATEGORIES = {"hombres": "Hombres", "mujeres": "Mujeres"}
CLUB_LINE_RE = re.compile(r"Liga:\s*(.+?)\s*\|.*?Club:\s*(.+)$", re.IGNORECASE)
STOPWORDS = {"el", "la", "los", "las", "de", "del", "y", "e"}
FUZZY_CUTOFF = 85.0
MIN_SCORE = 0.6


@dataclass
class RawMatch:
    category: str
    court: int
    time: str
    home: str
    away: str


@dataclass
class Fixture:
    number: int
    day: date | None
    matches: list[RawMatch] = field(default_factory=list)


@dataclass
class TurnosResult:
    text: str
    fecha: int
    day: date | None
    matches: int
    warnings: list[str] = field(default_factory=list)
    aliases: dict[str, dict[str, str]] = field(default_factory=dict)

    def to_dict(self) -> dict:
        return {
            "fecha": self.fecha,
            "date": self.day.isoformat() if self.day else None,
            "matchCount": self.matches,
            "warnings": self.warnings,
            "aliases": self.aliases,
        }


def clubs_from_page_texts(page_texts: list[str]) -> dict[str, list[str]]:
    """Nombres reales de los clubes del masivo, por rama (Hombres/Mujeres)."""
    clubs: dict[str, list[str]] = {"Hombres": [], "Mujeres": []}
    for text in page_texts:
        for raw in (text or "").splitlines():
            line = raw.replace("\xa0", " ")
            hit = CLUB_LINE_RE.search(line)
            if not hit:
                continue
            liga = hit.group(1).strip().lower()
            name = re.sub(r"\s+", " ", hit.group(2)).strip(" |")
            category = "Mujeres" if liga.startswith("mujeres") else "Hombres"
            if name and name not in clubs[category]:
                clubs[category].append(name)
    return clubs


# ---------------------------------------------------------------- Excel

def _is_time(value: object) -> bool:
    return isinstance(value, (time, datetime))


def _format_time(value: time | datetime) -> str:
    return f"{value.hour:02d}:{value.minute:02d}"


def _label_number(value: object) -> int | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    return None


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _read_sheet(sheet, category: str) -> dict[int, Fixture]:
    """Una Fixture por tabla numerada (se ignoran X, XX y PROMO)."""
    rows = list(sheet.iter_rows(values_only=True))
    fixtures: dict[int, Fixture] = {}
    for start, row in enumerate(rows):
        number = _label_number(row[0] if row else None)
        second = row[1] if len(row) > 1 else None
        if number is None or not (isinstance(second, str) and "hora" in second.lower()):
            continue

        # cada cancha usa la última columna de hora que tiene a su izquierda
        courts: dict[int, tuple[int, int]] = {}
        time_col = 1
        for col, header in enumerate(row):
            if not isinstance(header, str):
                continue
            label = header.strip().lower()
            if label.startswith("hora"):
                time_col = col
                continue
            hit = re.match(r"cancha\s+(\d+)$", label)
            if hit:
                courts[col] = (int(hit.group(1)), time_col)

        fixture = Fixture(number=number, day=None)
        for data in rows[start + 1 :]:
            hour = data[1] if len(data) > 1 else None
            if not _is_time(hour):
                break
            if fixture.day is None:
                fixture.day = _as_date(data[0])
            for col, (court, hour_col) in courts.items():
                cell = data[col] if col < len(data) else None
                if not isinstance(cell, str) or not cell.strip():
                    continue
                slot = data[hour_col] if hour_col < len(data) else None
                if not _is_time(slot):
                    continue
                parts = [part.strip() for part in cell.split("-") if part.strip()]
                if len(parts) != 2:
                    fixture.matches.append(
                        RawMatch(category, court, _format_time(slot), cell.strip(), "")
                    )
                    continue
                fixture.matches.append(
                    RawMatch(category, court, _format_time(slot), parts[0], parts[1])
                )
        fixtures[number] = fixture
    return fixtures


def read_fixture(xlsx_bytes: bytes) -> tuple[int, date | None, list[RawMatch]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=False)
    except Exception as error:  # noqa: BLE001
        raise ValueError("No pude abrir el Excel de turnos y canchas.") from error

    by_category: dict[str, dict[int, Fixture]] = {}
    for sheet in workbook.worksheets:
        category = SHEET_CATEGORIES.get(sheet.title.strip().lower())
        if category:
            by_category[category] = _read_sheet(sheet, category)
    if not by_category:
        raise ValueError("El Excel no tiene las hojas HOMBRES y MUJERES.")

    with_matches = [
        {number for number, fixture in tables.items() if fixture.matches}
        for tables in by_category.values()
    ]
    common = set.intersection(*with_matches) if with_matches else set()
    candidates = common or set.union(*with_matches)
    if not candidates:
        raise ValueError("No encontré partidos cargados en el Excel.")
    number = max(candidates)

    matches: list[RawMatch] = []
    day: date | None = None
    for tables in by_category.values():
        fixture = tables.get(number)
        if fixture:
            matches.extend(fixture.matches)
            day = day or fixture.day
    return number, day, matches


# ------------------------------------------------------------- cruce de nombres

def _tokens(name: str, *, drop_stopwords: bool) -> list[str]:
    words = normalize_name(name).split()
    if drop_stopwords:
        words = [word for word in words if word not in STOPWORDS]
    return words


def _squash(name: str) -> str:
    return "".join(_tokens(name, drop_stopwords=True))


def _score(abbr: str, club: str) -> float:
    """0..1 de qué tan bien la abreviatura describe al club."""
    abbr_tokens = _tokens(abbr, drop_stopwords=True)
    club_tokens = _tokens(club, drop_stopwords=True)
    if not abbr_tokens or not club_tokens:
        return 0.0
    abbr_squash = "".join(abbr_tokens)

    if abbr_squash == "".join(club_tokens):
        return 1.0

    if len(abbr_squash) >= 3:
        for with_stopwords in (True, False):
            initials = "".join(
                word[0] for word in _tokens(club, drop_stopwords=not with_stopwords)
            )
            if initials == abbr_squash:
                return 0.95

    # cada palabra de la abreviatura es el comienzo de una palabra distinta del club
    free = list(club_tokens)
    covered = 0
    for token in abbr_tokens:
        found = next((word for word in free if word.startswith(token)), None)
        if found is None:
            covered = -1
            break
        free.remove(found)
        covered += 1
    if covered > 0:
        return 0.6 + 0.4 * covered / len(club_tokens)

    ratio = fuzz.ratio(abbr_squash, "".join(club_tokens))
    if ratio >= FUZZY_CUTOFF:
        return 0.6 * ratio / 100.0
    return 0.0


def resolve_abbreviations(
    abbreviations: list[str],
    clubs: list[str],
    manual: dict[str, str] | None = None,
) -> dict[str, str | None]:
    """abreviatura -> nombre real (o None si no se pudo identificar con seguridad)."""
    manual = {normalize_name(key): value for key, value in (manual or {}).items()}
    result: dict[str, str | None] = {}
    pending: list[str] = []
    taken: set[str] = set()

    for abbr in dict.fromkeys(abbreviations):
        forced = manual.get(normalize_name(abbr))
        if forced:
            result[abbr] = forced
            taken.add(forced)
        else:
            pending.append(abbr)

    pairs = sorted(
        (
            (_score(abbr, club), abbr, club)
            for abbr in pending
            for club in clubs
            if club not in taken
        ),
        key=lambda item: -item[0],
    )
    for abbr in pending:
        result[abbr] = None

    ambiguous: set[str] = set()
    for index, (score, abbr, club) in enumerate(pairs):
        if score < MIN_SCORE:
            break
        if result[abbr] is not None or abbr in ambiguous or club in taken:
            continue
        rival = next(
            (
                other
                for other_score, other_abbr, other in pairs[index + 1 :]
                if other_abbr == abbr and other not in taken and other_score == score
            ),
            None,
        )
        if rival is not None:
            ambiguous.add(abbr)
            continue
        result[abbr] = club
        taken.add(club)
    return result


def load_manual_aliases(paths: list[Path]) -> dict[str, dict[str, str]]:
    """public/alias-equipos.json: {"Hombres": {"abreviatura": "Nombre real"}, "Mujeres": {...}}"""
    merged: dict[str, dict[str, str]] = {"Hombres": {}, "Mujeres": {}}
    for path in paths:
        if not path.exists():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if not isinstance(data, dict):
            continue
        for category, table in data.items():
            key = SHEET_CATEGORIES.get(str(category).strip().lower())
            if key and isinstance(table, dict):
                merged[key].update({str(k): str(v) for k, v in table.items()})
    return merged


# ------------------------------------------------------------------ armado

def build_schedule_text(
    xlsx_bytes: bytes,
    clubs: dict[str, list[str]],
    manual_aliases: dict[str, dict[str, str]] | None = None,
) -> TurnosResult:
    """Excel + clubes del masivo -> texto en el formato que ya entiende `parse_schedule`."""
    number, day, raw_matches = read_fixture(xlsx_bytes)
    manual_aliases = manual_aliases or {}
    warnings: list[str] = []
    used_aliases: dict[str, dict[str, str]] = {}
    resolved: list[RawMatch] = []

    for category in ("Hombres", "Mujeres"):
        category_matches = [match for match in raw_matches if match.category == category]
        if not category_matches:
            continue
        abbreviations = [
            name for match in category_matches for name in (match.home, match.away) if name
        ]
        if not clubs.get(category):
            warnings.append(
                f"El masivo no tiene clubes de {category}: no puedo identificar los equipos."
            )
            continue
        names = resolve_abbreviations(
            abbreviations, clubs[category], manual_aliases.get(category)
        )
        used_aliases[category] = {abbr: club for abbr, club in names.items() if club}
        for match in category_matches:
            home, away = names.get(match.home), names.get(match.away)
            if not match.away or not home or not away:
                unknown = [
                    f"«{raw}»"
                    for raw, full in ((match.home, home), (match.away, away))
                    if raw and not full
                ] or [f"«{match.home}»"]
                warnings.append(
                    f"No identifiqué {', '.join(unknown)} ({category}, cancha {match.court}, "
                    f"{match.time}). Agregalo a alias-equipos.json."
                )
                continue
            resolved.append(RawMatch(category, match.court, match.time, home, away))

    lines: list[str] = []
    for category in ("Hombres", "Mujeres"):
        group = [match for match in resolved if match.category == category]
        if not group:
            continue
        lines += [f"{category}:", ""]
        for court in sorted({match.court for match in group}):
            lines.append(f"Cancha {court}")
            for match in sorted(
                (item for item in group if item.court == court), key=lambda item: item.time
            ):
                lines.append(f"{match.time}: {match.home} vs {match.away}")
            lines.append("")
    return TurnosResult(
        text="\n".join(lines).strip() + "\n",
        fecha=number,
        day=day,
        matches=len(resolved),
        warnings=warnings,
        aliases=used_aliases,
    )
